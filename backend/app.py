from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
import requests
import tempfile
import os
import json
import PyPDF2
import docx
import openpyxl

app = FastAPI()

# Allow frontend to talk to backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = tempfile.gettempdir()
uploaded_file_path = None
document_content = ""   # parsed file text
chat_history = []       # store messages


# ---------------- Parse File ---------------- #
def parse_file(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    if ext == ".pdf":
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() or ""

    elif ext in [".docx"]:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"

    elif ext in [".xlsx", ".xlsm"]:
        wb = openpyxl.load_workbook(file_path)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) if cell else "" for cell in row])
                text += row_text + "\n"

    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

    return text.strip()


# ---------------- Upload File ---------------- #
@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    global uploaded_file_path, document_content, chat_history
    file_location = os.path.join(UPLOAD_DIR, file.filename)

    with open(file_location, "wb") as f:
        f.write(await file.read())

    uploaded_file_path = file_location
    chat_history = []  # reset chat when new file is uploaded
    document_content = parse_file(uploaded_file_path)

    return {"message": f"File '{file.filename}' uploaded successfully"}


# ---------------- Ask Question ---------------- #
@app.post("/ask")
async def ask_question(question: str = Form(...)):
    global document_content, chat_history

    if not document_content:
        return {"answer": "Please upload a document first."}

    # Append user message
    chat_history.append({"role": "user", "content": question})

    # Build messages: system + document + chat history
    messages = [
        {"role": "system", "content": "You are a helpful assistant. Only use the document for answers."},
        {"role": "system", "content": f"Document:\n{document_content}"}
    ] + chat_history

    payload = {
        "model": "llama3.2:latest",  # from `ollama list`
        "messages": messages
    }

    response = requests.post("http://localhost:11434/api/chat", json=payload, stream=True)

    if response.status_code == 200:
        final_answer = ""
        for line in response.iter_lines():
            if line:
                try:
                    obj = json.loads(line)
                    content = obj.get("message", {}).get("content")
                    if content:
                        final_answer += content
                except json.JSONDecodeError:
                    continue
        
        # Append assistant answer to history
        chat_history.append({"role": "assistant", "content": final_answer})

        return {"answer": final_answer}
    else:
        return {"answer": "Error contacting Ollama"}
